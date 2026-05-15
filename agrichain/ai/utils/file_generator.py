import os
from pathlib import Path


def generate_docx(result: dict, output_path: str) -> str:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    title = doc.add_heading("AgriChain - Weekly Farm Plan", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Farmer Information", level=1)
    info_table = doc.add_table(rows=6, cols=2)
    info_table.style = "Light Shading Accent 1"
    data = [
        ("Farmer Name", result.get("farmer_name", "")),
        ("Location", result.get("location", "")),
        ("Crop", result.get("crop", "")),
        ("Farm Size", f'{result.get("farm_size", "")} hectares'),
        ("Soil Type", result.get("soil_type", "N/A")),
        ("Fertilization Method", result.get("fertilization_method", "N/A")),
    ]
    for i, (key, val) in enumerate(data):
        info_table.cell(i, 0).text = key
        info_table.cell(i, 1).text = str(val)

    plan_text = result.get("farm_plan", "No plan generated.")
    doc.add_heading("Farm Plan", level=1)

    for line in plan_text.strip().split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("1.") or line.startswith("2.") or line.startswith("3.") or (
            "WHAT TO DO" in line.upper() or "WHEN AND WHERE" in line.upper()
            or "FINANCING" in line.upper() or "FERTILIZATION" in line.upper()
        ):
            doc.add_heading(line, level=2)
        else:
            p = doc.add_paragraph(line)

    agent_reports = result.get("agent_reports", {})
    doc.add_heading("Agent Reports", level=1)
    for agent_name, report_text in agent_reports.items():
        doc.add_heading(agent_name.title(), level=2)
        for line in report_text.strip().split("\n"):
            doc.add_paragraph(line.strip())

    execution_times = result.get("execution_times", {})
    doc.add_heading("Execution Times", level=1)
    times_table = doc.add_table(rows=len(execution_times), cols=2)
    times_table.style = "Light Shading Accent 1"
    for i, (key, val) in enumerate(execution_times.items()):
        times_table.cell(i, 0).text = key.title()
        times_table.cell(i, 1).text = str(val)

    doc.save(output_path)
    return output_path


def generate_pdf(result: dict, output_path: str) -> str:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "AgriChain - Weekly Farm Plan", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Farmer Information", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    info_lines = [
        ("Farmer Name", result.get("farmer_name", "")),
        ("Location", result.get("location", "")),
        ("Crop", result.get("crop", "")),
        ("Farm Size", f'{result.get("farm_size", "")} hectares'),
        ("Soil Type", result.get("soil_type", "N/A")),
        ("Fertilization Method", result.get("fertilization_method", "N/A")),
    ]
    for key, val in info_lines:
        pdf.cell(0, 6, f"{key}: {val}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Farm Plan", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)

    plan_text = result.get("farm_plan", "No plan generated.")
    for line in plan_text.strip().split("\n"):
        line = line.strip()
        if not line:
            continue
        if (line.startswith("1.") or line.startswith("2.") or line.startswith("3.")
                or "WHAT TO DO" in line.upper() or "WHEN AND WHERE" in line.upper()
                or "FINANCING" in line.upper() or "FERTILIZATION" in line.upper()):
            pdf.set_font("Helvetica", "B", 10)
            pdf.multi_cell(0, 6, line)
            pdf.set_font("Helvetica", "", 10)
        else:
            pdf.multi_cell(0, 6, line)
    pdf.ln(5)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Agent Reports", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    agent_reports = result.get("agent_reports", {})
    for agent_name, report_text in agent_reports.items():
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 6, agent_name.title(), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        for line in report_text.strip().split("\n"):
            pdf.multi_cell(0, 5, line.strip())

    execution_times = result.get("execution_times", {})
    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Execution Times", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    for key, val in execution_times.items():
        pdf.cell(0, 6, f"{key.title()}: {val}", new_x="LMARGIN", new_y="NEXT")

    pdf.output(output_path)
    return output_path


def save_farm_plan_files(result: dict, output_dir: str = None) -> dict:
    if output_dir is None:
        output_dir = str(Path.cwd() / "farm_plans")
    os.makedirs(output_dir, exist_ok=True)

    farmer_name = result.get("farmer_name", "farmer").replace(" ", "_")
    crop = result.get("crop", "crop").replace(" ", "_")
    base = f"{farmer_name}_{crop}"

    docx_path = os.path.join(output_dir, f"{base}.docx")
    pdf_path = os.path.join(output_dir, f"{base}.pdf")

    generate_docx(result, docx_path)
    generate_pdf(result, pdf_path)

    return {"docx": docx_path, "pdf": pdf_path}
